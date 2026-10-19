from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
OUT_PATH = OUT_DIR / "Invenia_AI_Commercial_Exhibition_Deck.pptx"
LOGO = ROOT / "frontend" / "src" / "assets" / "Invenia India Logo.png"
SHOT_DIR = ROOT / "docs" / "site" / "assets" / "screenshots"

SLIDE_W = 13.333
SLIDE_H = 7.5

NAVY = "071B3D"
NAVY_2 = "0D2D61"
BLUE = "2563EB"
BLUE_2 = "3B82F6"
CYAN = "22D3EE"
TEAL = "10B981"
AMBER = "F59E0B"
CORAL = "F43F5E"
VIOLET = "7C3AED"
WHITE = "FFFFFF"
INK = "0F172A"
SLATE = "475569"
MUTED = "64748B"
LIGHT = "F4F8FF"
PALE_BLUE = "EAF2FF"
PALE_CYAN = "E8FBFF"
PALE_GREEN = "EAFBF4"
PALE_AMBER = "FFF8E7"
LINE = "D8E6F8"

TITLE_FONT = "Aptos Display"
BODY_FONT = "Aptos"


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.replace("#", "")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def add_shape(slide, shape_type, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_rect(slide, x, y, w, h, fill, line=None, rounded=False):
    return add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line)


def add_circle(slide, x, y, d, fill, line=None):
    return add_shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill, line)


def set_shape_transparency(shape, transparency: int):
    shape.fill.transparency = max(0, min(100, transparency))


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 18,
    color: str = INK,
    bold: bool = False,
    font: str = BODY_FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.02,
    line_spacing: float | None = None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    if line_spacing is not None:
        paragraph.line_spacing = line_spacing
    run = paragraph.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, segments, x, y, w, h, size=18, color=INK, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for item in segments:
        run = paragraph.add_run()
        run.text = item[0]
        run.font.name = item[4] if len(item) > 4 else BODY_FONT
        run.font.size = Pt(item[1] if len(item) > 1 else size)
        run.font.bold = item[2] if len(item) > 2 else False
        run.font.color.rgb = rgb(item[3] if len(item) > 3 else color)
    return box


def add_bullets(slide, items: Sequence[str], x, y, w, h, size=15, color=SLATE, bullet_color=BLUE, gap=5):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = f"{chr(8226)}  {item}"
        p.font.name = BODY_FONT
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
    return box


def add_badge(slide, text, x, y, w, fill=PALE_BLUE, color=BLUE, line=None, size=10):
    add_rect(slide, x, y, w, 0.32, fill, line or fill, rounded=True)
    add_text(slide, text.upper(), x + 0.08, y + 0.035, w - 0.16, 0.22, size, color, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_icon_badge(slide, label, x, y, d=0.52, fill=BLUE, text_color=WHITE, size=12):
    add_circle(slide, x, y, d, fill)
    add_text(slide, label, x, y + 0.02, d, d - 0.04, size, text_color, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_logo(slide, x=11.95, y=0.23, w=0.82, on_dark=False):
    if on_dark:
        pill = add_rect(slide, x - 0.09, y - 0.04, w + 0.18, 0.58, WHITE, WHITE, rounded=True)
        set_shape_transparency(pill, 3)
    slide.shapes.add_picture(str(LOGO), Inches(x), Inches(y), width=Inches(w), height=Inches(w * 0.724))


def add_slide_number(slide, number, dark=False):
    color = "AFC8EF" if dark else MUTED
    add_text(slide, f"{number:02d}", 12.42, 7.08, 0.55, 0.22, 9, color, True, align=PP_ALIGN.RIGHT)


def add_footer(slide, number, dark=False):
    color = "AFC8EF" if dark else MUTED
    line_color = "244879" if dark else LINE
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(7.03), Inches(12.82), Inches(7.03))
    line.line.color.rgb = rgb(line_color)
    line.line.width = Pt(0.8)
    add_text(slide, "INVENIA AI  |  Intelligent Recruitment and Assessment", 0.52, 7.08, 5.7, 0.2, 8, color, True)
    add_slide_number(slide, number, dark)


def add_background(slide, dark=False, accent=True):
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, NAVY if dark else LIGHT)
    if dark:
        glow = add_circle(slide, 9.7, -1.5, 5.0, BLUE)
        set_shape_transparency(glow, 68)
        glow2 = add_circle(slide, -1.8, 5.4, 3.4, CYAN)
        set_shape_transparency(glow2, 83)
    elif accent:
        band = add_rect(slide, 0, 0, 0.11, SLIDE_H, BLUE)
        set_shape_transparency(band, 0)
        glow = add_circle(slide, 10.8, -1.7, 4.2, PALE_CYAN)
        set_shape_transparency(glow, 8)


def add_title(slide, kicker, title, subtitle=None, dark=False, number=None):
    color = WHITE if dark else INK
    muted = "C4D8F7" if dark else MUTED
    add_text(slide, kicker.upper(), 0.62, 0.38, 6.8, 0.25, 10, CYAN if dark else BLUE, True)
    add_text(slide, title, 0.62, 0.73, 11.2, 0.55, 28, color, True, TITLE_FONT)
    if subtitle:
        add_text(slide, subtitle, 0.62, 1.34, 11.1, 0.48, 12.5, muted)
    add_logo(slide, on_dark=dark)
    if number is not None:
        add_footer(slide, number, dark)


def add_metric(slide, value, label, x, y, w=1.85, fill=WHITE, value_color=INK, label_color=MUTED, line=LINE, note=None):
    add_rect(slide, x, y, w, 1.0, fill, line, rounded=True)
    add_text(slide, value, x + 0.16, y + 0.14, w - 0.32, 0.42, 24, value_color, True, TITLE_FONT)
    add_text(slide, label, x + 0.16, y + 0.59, w - 0.32, 0.22, 9.5, label_color, True)
    if note:
        add_text(slide, note, x + 0.16, y + 0.80, w - 0.32, 0.15, 7.5, label_color)


def add_card(slide, title, body, x, y, w, h, icon=None, accent=BLUE, fill=WHITE, line=LINE, title_size=15, body_size=10.5):
    add_rect(slide, x, y, w, h, fill, line, rounded=True)
    if icon:
        add_icon_badge(slide, icon, x + 0.18, y + 0.18, 0.48, accent, WHITE, 10)
        title_x = x + 0.78
        title_w = w - 0.96
    else:
        title_x = x + 0.22
        title_w = w - 0.44
    add_text(slide, title, title_x, y + 0.18, title_w, 0.34, title_size, INK, True, TITLE_FONT)
    add_text(slide, body, x + 0.22, y + 0.67, w - 0.44, h - 0.82, body_size, SLATE)


def add_picture_contain(slide, path: Path, x, y, w, h, line=LINE, rounded_frame=True):
    if rounded_frame:
        add_rect(slide, x - 0.08, y - 0.08, w + 0.16, h + 0.16, WHITE, line, rounded=True)
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(w / iw, h / ih)
    pw = iw * scale
    ph = ih * scale
    px = x + (w - pw) / 2
    py = y + (h - ph) / 2
    slide.shapes.add_picture(str(path), Inches(px), Inches(py), width=Inches(pw), height=Inches(ph))


def add_picture_crop(slide, path: Path, x, y, w, h, line=LINE):
    add_rect(slide, x - 0.08, y - 0.08, w + 0.16, h + 0.16, WHITE, line, rounded=True)
    with Image.open(path) as image:
        iw, ih = image.size
    target_ratio = w / h
    image_ratio = iw / ih
    shape = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    if image_ratio > target_ratio:
        visible_ratio = target_ratio / image_ratio
        crop = (1 - visible_ratio) / 2
        shape.crop_left = crop
        shape.crop_right = crop
    elif image_ratio < target_ratio:
        visible_ratio = image_ratio / target_ratio
        crop = (1 - visible_ratio) / 2
        shape.crop_top = crop
        shape.crop_bottom = crop
    return shape


def add_line(slide, x1, y1, x2, y2, color=LINE, width=1.5, dash=False):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    if dash:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return line


def add_process_step(slide, number, title, body, x, y, w, accent=BLUE, dark=False):
    fill = "102D5E" if dark else WHITE
    line = "29548C" if dark else LINE
    title_color = WHITE if dark else INK
    body_color = "C4D8F7" if dark else SLATE
    add_rect(slide, x, y, w, 1.2, fill, line, rounded=True)
    add_icon_badge(slide, str(number), x + 0.16, y + 0.2, 0.48, accent, WHITE, 11)
    add_text(slide, title, x + 0.78, y + 0.17, w - 0.95, 0.31, 13.5, title_color, True, TITLE_FONT)
    add_text(slide, body, x + 0.18, y + 0.65, w - 0.36, 0.38, 9.5, body_color)


def add_weight_bar(slide, label, score, weight, contribution, x, y, w, color):
    add_text(slide, label, x, y, 2.25, 0.22, 10.5, INK, True)
    add_text(slide, f"{score}% x {weight}%", x + 2.2, y, 1.35, 0.22, 9.5, MUTED, align=PP_ALIGN.RIGHT)
    add_rect(slide, x, y + 0.32, w, 0.16, "E5EDF8", "E5EDF8", rounded=True)
    add_rect(slide, x, y + 0.32, w * score / 100, 0.16, color, color, rounded=True)
    add_text(slide, f"{contribution:.1f} pts", x + w + 0.12, y + 0.22, 0.72, 0.28, 9.5, color, True, align=PP_ALIGN.RIGHT)


def add_formula_box(slide, formula, x, y, w, h, fill=NAVY, accent=CYAN):
    add_rect(slide, x, y, w, h, fill, fill, rounded=True)
    add_rect(slide, x, y, 0.08, h, accent, accent, rounded=True)
    add_text(slide, formula, x + 0.25, y + 0.12, w - 0.45, h - 0.22, 15, WHITE, True, font="Aptos Mono", valign=MSO_ANCHOR.MIDDLE)


def create_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "Invenia AI - Commercial Technology Exhibition Deck"
    prs.core_properties.subject = "AI-powered recruitment, assessment, interview, and analytics platform"
    prs.core_properties.author = "Invenia AI"
    prs.core_properties.company = "Invenia AI"
    prs.core_properties.comments = "Generated from the implemented Invenia AI product architecture and live Railway deployment."
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank)
    add_background(slide, dark=True)
    add_rect(slide, 0.62, 0.58, 1.32, 0.88, WHITE, WHITE, rounded=True)
    slide.shapes.add_picture(str(LOGO), Inches(0.72), Inches(0.68), width=Inches(1.12), height=Inches(0.81))
    add_badge(slide, "Technology Exhibition Edition", 2.18, 0.83, 2.55, "173A70", CYAN, "315990", 9)
    add_text(slide, "INVENIA AI", 0.7, 2.03, 8.2, 0.82, 42, WHITE, True, TITLE_FONT)
    add_text(slide, "The intelligent recruitment and assessment operating system", 0.72, 2.92, 9.6, 0.52, 22, "BFD8FF", True, TITLE_FONT)
    add_text(slide, "From job design and CV intelligence to secure assessment, voice interview, evidence-grounded scoring, and recruiter decision support.", 0.72, 3.65, 8.55, 0.92, 15, "DCEAFF")
    add_rect(slide, 9.65, 1.52, 2.8, 3.7, "102D5E", "29548C", rounded=True)
    for index, (label, value, color) in enumerate([
        ("ROLES", "2", CYAN),
        ("APPLICATION STAGES", "5", TEAL),
        ("AI INTERVIEW TURNS", "3", AMBER),
        ("DEPLOYMENT", "LIVE", BLUE_2),
    ]):
        yy = 1.82 + index * 0.82
        add_text(slide, label, 9.95, yy, 1.55, 0.2, 8.5, "AFC8EF", True)
        add_text(slide, value, 11.3, yy - 0.08, 0.8, 0.37, 19, color, True, TITLE_FONT, align=PP_ALIGN.RIGHT)
        if index < 3:
            add_line(slide, 9.95, yy + 0.42, 12.12, yy + 0.42, "29548C", 0.7)
    add_text(slide, "Commercial product showcase", 0.72, 6.33, 4.3, 0.25, 10, CYAN, True)
    live = add_text(slide, "frontend-production-fad4.up.railway.app", 0.72, 6.65, 5.1, 0.24, 10, "AFC8EF")
    live.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://frontend-production-fad4.up.railway.app/"
    add_footer(slide, 1, dark=True)

    # 2. Executive snapshot
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Product introduction", "A single decision layer for modern hiring", "Invenia AI connects the candidate journey, recruiter workflow, assessments, media, and analytics in one governed platform.", number=2)
    add_picture_crop(slide, SHOT_DIR / "recruiter-overview.png", 0.68, 2.02, 7.25, 4.55)
    x = 8.28
    add_metric(slide, "2", "ROLE-SPECIFIC PORTALS", x, 2.02, 2.05, PALE_BLUE, BLUE)
    add_metric(slide, "5", "APPLICATION STAGES", x + 2.25, 2.02, 2.05, PALE_GREEN, TEAL)
    add_metric(slide, "69", "BACKEND API ROUTES", x, 3.25, 2.05, PALE_CYAN, NAVY_2)
    add_metric(slide, "LIVE", "RAILWAY DEPLOYMENT", x + 2.25, 3.25, 2.05, PALE_AMBER, AMBER)
    add_card(slide, "Product thesis", "Replace fragmented hiring tools with one evidence-driven workflow where every score is traceable to candidate evidence, enabled rounds, and recruiter-approved decisions.", x, 4.56, 4.3, 1.98, "AI", BLUE, WHITE, LINE, 16, 11)

    # 3. Problem statement
    slide = prs.slides.add_slide(blank)
    add_background(slide, dark=True)
    add_title(slide, "Problem statement", "Hiring is fragmented, slow, and difficult to trust", "Enterprises often combine disconnected ATS, assessment, interview, file, and reporting tools - creating delays and weak decision traceability.", dark=True, number=3)
    problems = [
        ("P1", "Manual screening overload", "Recruiters compare CVs against complex job requirements by hand, creating inconsistent shortlists.", CORAL),
        ("P2", "Generic scores", "Keyword-only scores do not explain evidence, job context, or why one candidate outranks another.", AMBER),
        ("P3", "Disconnected rounds", "Aptitude, technical, coding, and interviews live in separate products and data silos.", VIOLET),
        ("P4", "Interview integrity", "Remote interviews require identity checks, camera evidence, transcripts, and auditable behavior signals.", CYAN),
        ("P5", "Poor recruiter visibility", "Decision makers need job-wise comparisons, recordings, reports, and exports in one place.", TEAL),
        ("P6", "Candidate uncertainty", "Applicants need clear status, stage gates, notifications, and transparent next steps.", BLUE_2),
    ]
    for i, (icon, title, body, accent) in enumerate(problems):
        col = i % 3
        row = i // 3
        x = 0.7 + col * 4.18
        y = 2.05 + row * 2.08
        add_rect(slide, x, y, 3.82, 1.68, "102D5E", "244879", rounded=True)
        add_icon_badge(slide, icon, x + 0.18, y + 0.18, 0.5, accent, WHITE, 9)
        add_text(slide, title, x + 0.82, y + 0.18, 2.75, 0.3, 14, WHITE, True, TITLE_FONT)
        add_text(slide, body, x + 0.2, y + 0.72, 3.4, 0.72, 10, "C4D8F7")

    # 4. Objectives
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Product objectives", "Designed around measurable hiring outcomes", "The implementation prioritizes decision quality, candidate clarity, assessment integrity, and operational speed.", number=4)
    objectives = [
        ("01", "Unify", "One platform from authentication to final hiring analysis", BLUE),
        ("02", "Automate", "Structured job and profile extraction plus role-specific question generation", VIOLET),
        ("03", "Explain", "Evidence-grounded Job ATS with transparent mathematical components", CYAN),
        ("04", "Control", "Recruiter approval gates, optional rounds, immutable decisions, account controls", TEAL),
        ("05", "Secure", "Fullscreen monitoring, camera, mic test, snapshots, recording, and warnings", CORAL),
        ("06", "Compare", "Job-wise candidate ranking, detailed reports, filters, and exports", AMBER),
    ]
    for i, (num, title, body, accent) in enumerate(objectives):
        col = i % 3
        row = i // 3
        add_card(slide, title, body, 0.68 + col * 4.18, 2.05 + row * 2.05, 3.82, 1.65, num, accent, WHITE, LINE, 16, 10.5)
    add_formula_box(slide, "SUCCESS = higher-quality evidence + faster screening + controlled assessments + explainable decisions", 0.7, 6.22, 11.95, 0.58, NAVY, CYAN)

    # 5. Role ecosystem
    slide = prs.slides.add_slide(blank)
    add_background(slide, dark=True)
    add_title(slide, "Experience architecture", "Two portals. One synchronized hiring record.", "Candidates control their career journey while recruiters control screening, stage access, analysis, and final decisions.", dark=True, number=5)
    add_rect(slide, 0.72, 2.0, 3.35, 4.15, "102D5E", "2A5B98", rounded=True)
    add_icon_badge(slide, "C", 1.05, 2.32, 0.78, BLUE_2, WHITE, 20)
    add_text(slide, "Candidate portal", 1.98, 2.36, 1.65, 0.35, 20, WHITE, True, TITLE_FONT)
    add_bullets(slide, [
        "Profile and CV autofill",
        "Profile ATS and job-specific ATS",
        "Job search, save, apply, revoke",
        "Approval-gated assessment access",
        "Secure voice interview and status tracking",
    ], 1.0, 3.15, 2.8, 2.25, 12, "DCEAFF")
    add_rect(slide, 9.25, 2.0, 3.35, 4.15, "102D5E", "2A5B98", rounded=True)
    add_icon_badge(slide, "R", 9.58, 2.32, 0.78, VIOLET, WHITE, 20)
    add_text(slide, "Recruiter portal", 10.5, 2.36, 1.75, 0.35, 20, WHITE, True, TITLE_FONT)
    add_bullets(slide, [
        "Smart job creation and stage rules",
        "Job-wise candidate comparison",
        "Approve, reject, activate, deactivate, delete",
        "Interview media, transcript, and reports",
        "Analytics, rankings, and spreadsheet exports",
    ], 9.53, 3.15, 2.78, 2.25, 12, "DCEAFF")
    add_rect(slide, 4.55, 2.42, 4.25, 3.35, "173A70", "315990", rounded=True)
    add_text(slide, "SHARED INTELLIGENCE CORE", 4.95, 2.76, 3.45, 0.3, 11, CYAN, True, align=PP_ALIGN.CENTER)
    for idx, (label, accent) in enumerate([("Identity + roles", BLUE_2), ("Jobs + applications", TEAL), ("Workflow + scores", AMBER), ("Media + reports", CORAL), ("Analytics + exports", VIOLET)]):
        y = 3.28 + idx * 0.44
        add_circle(slide, 4.98, y + 0.03, 0.18, accent)
        add_text(slide, label, 5.28, y, 2.95, 0.24, 11, WHITE, True)
    add_line(slide, 4.08, 4.05, 4.55, 4.05, CYAN, 2)
    add_line(slide, 8.8, 4.05, 9.25, 4.05, CYAN, 2)

    # 6. Workflow
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "End-to-end workflow", "From job intent to recruiter-ready evidence", "Every transition is persisted, role-gated, and visible to the appropriate participant.", number=6)
    steps = [
        ("1", "Create job", "Autofill or manual details; configure optional and compulsory rounds", BLUE),
        ("2", "Apply + scan", "CV is compared against the exact job and sent for screening", CYAN),
        ("3", "Recruiter gate", "Job-wise ranking supports an approve or reject decision", VIOLET),
        ("4", "Assess", "Only enabled aptitude, technical, and coding stages appear", AMBER),
        ("5", "Interview", "Camera, mic test, fullscreen, three voice questions, recording", CORAL),
        ("6", "Decide", "Final score, evidence, media, report, comparison, and export", TEAL),
    ]
    for i, (num, title, body, accent) in enumerate(steps):
        row = i // 3
        col = i % 3
        x = 0.68 + col * 4.2
        y = 2.08 + row * 2.02
        add_process_step(slide, num, title, body, x, y, 3.82, accent)
        if col < 2:
            add_line(slide, x + 3.82, y + 0.6, x + 4.16, y + 0.6, accent, 2)
    add_formula_box(slide, "Recruiter configuration controls the path. Candidate eligibility controls access. Persisted evidence controls the decision.", 0.7, 6.22, 11.92, 0.56, NAVY, TEAL)

    # 7. Candidate experience
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Candidate experience", "Clarity before, during, and after every application", "The portal combines opportunity discovery, profile intelligence, CV access, status tracking, and assessment readiness.", number=7)
    add_picture_crop(slide, SHOT_DIR / "candidate-jobs.png", 0.66, 1.94, 6.05, 3.82)
    add_picture_crop(slide, SHOT_DIR / "candidate-profile.png", 6.92, 1.94, 5.75, 3.82)
    labels = [
        ("DISCOVER", "Grid/list jobs, search, save, view, apply", 0.72, BLUE),
        ("UNDERSTAND", "Profile strength, CV preview, ATS insights", 3.75, CYAN),
        ("TRACK", "Applied status, stage lock, next-round access", 6.95, VIOLET),
        ("CONTROL", "Revoke once, avoid duplicates, receive notifications", 10.0, TEAL),
    ]
    for title, body, x, accent in labels:
        add_rect(slide, x, 6.02, 2.58, 0.72, WHITE, LINE, rounded=True)
        add_rect(slide, x, 6.02, 0.07, 0.72, accent, accent, rounded=True)
        add_text(slide, title, x + 0.18, 6.11, 1.05, 0.18, 8.5, accent, True)
        add_text(slide, body, x + 0.18, 6.34, 2.2, 0.24, 8.5, SLATE)

    # 8. Recruiter command center
    slide = prs.slides.add_slide(blank)
    add_background(slide, dark=True)
    add_title(slide, "Recruiter command center", "One view from workload to hiring intelligence", "Overview KPIs surface active jobs, total candidates, applications, interviews, hires, and time-to-hire alongside live pipeline data.", dark=True, number=8)
    add_picture_crop(slide, SHOT_DIR / "recruiter-overview.png", 0.75, 1.95, 8.3, 4.9, "244879")
    benefits = [
        ("01", "Pipeline visibility", "Recent applications and job demand in one dashboard", BLUE_2),
        ("02", "Job operations", "Create, edit, archive, export, and switch grid/list views", TEAL),
        ("03", "Candidate governance", "Approve/reject decisions and account lifecycle controls", VIOLET),
        ("04", "Decision intelligence", "Job ATS, final score, reports, recordings, comparisons", AMBER),
    ]
    for i, (num, title, body, accent) in enumerate(benefits):
        y = 2.05 + i * 1.15
        add_rect(slide, 9.4, y, 3.2, 0.92, "102D5E", "244879", rounded=True)
        add_icon_badge(slide, num, 9.58, y + 0.2, 0.42, accent, WHITE, 8.5)
        add_text(slide, title, 10.15, y + 0.15, 2.14, 0.24, 12, WHITE, True)
        add_text(slide, body, 10.15, y + 0.45, 2.12, 0.35, 8.8, "C4D8F7")

    # 9. Feature matrix
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Feature portfolio", "Commercial capability across the full hiring lifecycle", "Implemented functions are organized as reusable product modules rather than isolated demos.", number=9)
    groups = [
        ("ACCESS", ["Candidate sign-up/login", "Recruiter login", "JWT role gates", "Account controls"], BLUE),
        ("PROFILE", ["CV upload/preview", "Profile autofill", "ATS readiness", "Skills/preferences"], CYAN),
        ("JOBS", ["Smart autofill", "Grid/list views", "Currency/period", "Views/applicants"], TEAL),
        ("WORKFLOW", ["Approval gate", "Optional stages", "Duplicate prevention", "Application revoke"], VIOLET),
        ("ASSESSMENT", ["5 aptitude MCQs", "5 technical MCQs", "2 coding prompts", "Role-context generation"], AMBER),
        ("INTERVIEW", ["Face presence", "10s mic test", "3 voice questions", "Recording + snapshot"], CORAL),
        ("ANALYTICS", ["Job-wise ranking", "Detailed analysis", "Performance reports", "Top-N filters"], BLUE_2),
        ("OPERATIONS", ["Excel/CSV export", "CV one-click preview", "Media download", "Railway hosting"], NAVY_2),
    ]
    for i, (title, items, accent) in enumerate(groups):
        col = i % 4
        row = i // 4
        x = 0.62 + col * 3.13
        y = 1.95 + row * 2.32
        add_rect(slide, x, y, 2.85, 1.98, WHITE, LINE, rounded=True)
        add_text(slide, title, x + 0.2, y + 0.18, 1.9, 0.25, 11, accent, True)
        add_rect(slide, x + 2.38, y + 0.16, 0.27, 0.27, accent, accent, rounded=True)
        add_bullets(slide, items, x + 0.2, y + 0.58, 2.42, 1.15, 9.5, SLATE, accent, 3)
    add_text(slide, "Modular by design: capabilities can evolve independently without breaking the core application journey.", 0.68, 6.7, 11.8, 0.25, 10, MUTED, True, align=PP_ALIGN.CENTER)

    # 10. AI orchestration
    slide = prs.slides.add_slide(blank)
    add_background(slide, dark=True)
    add_title(slide, "Intelligence orchestration", "AI proposes. Deterministic systems validate and score.", "Invenia AI separates semantic reasoning from numeric calculation to reduce arbitrary scoring and improve auditability.", dark=True, number=10)
    add_rect(slide, 0.72, 2.0, 3.2, 4.25, "102D5E", "244879", rounded=True)
    add_text(slide, "INPUT SIGNALS", 1.0, 2.3, 2.65, 0.25, 11, CYAN, True, align=PP_ALIGN.CENTER)
    inputs = ["Job description", "CV and profile", "Assessment answers", "Interview transcript", "Proctoring events"]
    for i, item in enumerate(inputs):
        add_rect(slide, 1.02, 2.82 + i * 0.57, 2.6, 0.4, "173A70", "315990", rounded=True)
        add_text(slide, item, 1.18, 2.91 + i * 0.57, 2.28, 0.2, 10, WHITE, True)
    add_rect(slide, 4.42, 2.0, 4.5, 4.25, "173A70", "315990", rounded=True)
    add_text(slide, "INVENIA INTELLIGENCE LAYER", 4.82, 2.3, 3.7, 0.25, 11, CYAN, True, align=PP_ALIGN.CENTER)
    intelligence = [
        ("EXTRACT", "Structured profile and job fields", BLUE_2),
        ("CLASSIFY", "Matched, partial, or missing evidence", VIOLET),
        ("GENERATE", "MCQs, coding prompts, interview turns", AMBER),
        ("ANALYZE", "Interview performance and hiring risks", CORAL),
        ("VALIDATE", "Evidence quotes, schemas, confidence", TEAL),
    ]
    for i, (label, body, accent) in enumerate(intelligence):
        y = 2.77 + i * 0.62
        add_icon_badge(slide, str(i + 1), 4.8, y, 0.4, accent, WHITE, 8)
        add_text(slide, label, 5.36, y - 0.01, 1.0, 0.2, 9.5, accent, True)
        add_text(slide, body, 6.25, y - 0.01, 2.25, 0.28, 9.5, WHITE)
    add_rect(slide, 9.42, 2.0, 3.2, 4.25, "102D5E", "244879", rounded=True)
    add_text(slide, "GOVERNED OUTPUTS", 9.72, 2.3, 2.6, 0.25, 11, CYAN, True, align=PP_ALIGN.CENTER)
    outputs = ["Evidence-grounded Job ATS", "Versioned final score", "Question sets", "Interview report", "Recruiter recommendation"]
    for i, item in enumerate(outputs):
        add_rect(slide, 9.72, 2.82 + i * 0.57, 2.6, 0.4, "173A70", "315990", rounded=True)
        add_text(slide, item, 9.88, 2.91 + i * 0.57, 2.28, 0.2, 9.5, WHITE, True)
    add_line(slide, 3.92, 4.08, 4.42, 4.08, CYAN, 2)
    add_line(slide, 8.92, 4.08, 9.42, 4.08, CYAN, 2)

    # 11. Architecture
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "System architecture", "Cloud-hosted, API-driven, and modular", "Separated frontend and backend services connect to managed data, persistent media, and external intelligence services.", number=11)
    # Actors
    add_card(slide, "Candidate", "Profile, jobs, applications, assessments, interview", 0.55, 2.05, 2.1, 1.15, "C", BLUE, PALE_BLUE, "C7DBFA", 14, 9.5)
    add_card(slide, "Recruiter", "Jobs, pipeline, decisions, media, analytics", 0.55, 3.65, 2.1, 1.15, "R", VIOLET, "F3EEFF", "DED0FF", 14, 9.5)
    # Frontend
    add_rect(slide, 3.2, 2.35, 2.35, 2.15, NAVY, NAVY, rounded=True)
    add_text(slide, "FRONTEND", 3.53, 2.7, 1.7, 0.22, 10, CYAN, True, align=PP_ALIGN.CENTER)
    add_text(slide, "React 19 + Vite 8", 3.5, 3.1, 1.75, 0.3, 17, WHITE, True, TITLE_FONT, align=PP_ALIGN.CENTER)
    add_text(slide, "Tailwind CSS\nAxios + Router\nNginx container", 3.55, 3.54, 1.65, 0.72, 10, "DCEAFF", align=PP_ALIGN.CENTER)
    # Backend
    add_rect(slide, 6.03, 1.95, 2.7, 2.95, "173A70", "315990", rounded=True)
    add_text(slide, "APPLICATION API", 6.4, 2.28, 1.96, 0.22, 10, CYAN, True, align=PP_ALIGN.CENTER)
    add_text(slide, "FastAPI", 6.36, 2.68, 2.02, 0.36, 22, WHITE, True, TITLE_FONT, align=PP_ALIGN.CENTER)
    add_text(slide, "JWT + role gates\nWorkflow engine\nScoring + analytics\nUploads + media APIs", 6.42, 3.18, 1.9, 1.2, 10, "DCEAFF", align=PP_ALIGN.CENTER)
    # Data and AI
    targets = [
        ("PostgreSQL", "Users, jobs, workflows, reports", "DB", BLUE, 9.35, 1.85),
        ("Persistent volume", "CVs, snapshots, recordings", "FS", TEAL, 9.35, 3.15),
        ("Model APIs", "Extraction, evidence, questions, analysis", "AI", VIOLET, 9.35, 4.45),
        ("Voice service", "Interview question speech", "VO", AMBER, 9.35, 5.75),
    ]
    for title, body, icon, accent, x, y in targets:
        add_card(slide, title, body, x, y, 3.0, 0.92, icon, accent, WHITE, LINE, 13, 8.7)
    # Connectors
    add_line(slide, 2.65, 2.62, 3.2, 3.05, BLUE, 2)
    add_line(slide, 2.65, 4.22, 3.2, 3.82, VIOLET, 2)
    add_line(slide, 5.55, 3.42, 6.03, 3.42, CYAN, 2.5)
    for y in [2.31, 3.61, 4.91, 6.21]:
        add_line(slide, 8.73, 3.42, 9.35, y, NAVY_2, 1.5)
    add_badge(slide, "Railway - containerized production hosting", 3.2, 5.45, 5.55, PALE_CYAN, NAVY_2, "BDEAF3", 10)

    # 12. Tech stack
    slide = prs.slides.add_slide(blank)
    add_background(slide, dark=True)
    add_title(slide, "Technology bifurcation", "A modern full-stack implementation", "Each layer uses focused frameworks and services chosen for delivery speed, maintainability, and cloud portability.", dark=True, number=12)
    columns = [
        ("CLIENT", "UI", BLUE_2, ["React 19.2", "Vite 8", "React Router 7", "Tailwind CSS 3", "Heroicons", "Axios"]),
        ("SERVER", "API", CYAN, ["Python 3.11", "FastAPI 0.135", "Pydantic 2", "Uvicorn", "Requests", "Multipart uploads"]),
        ("DATA", "DB", TEAL, ["SQLAlchemy 2", "PostgreSQL 18", "SQLite local mode", "Persistent volumes", "JSON workflow state", "Binary media fields"]),
        ("INTELLIGENCE", "AI", VIOLET, ["Gemini 2.5 Flash", "Structured JSON", "Evidence validation", "Deterministic scoring", "ElevenLabs TTS", "Browser speech APIs"]),
        ("PLATFORM", "OPS", AMBER, ["Docker", "Railway", "Nginx", "CORS controls", "Health checks", "Environment secrets"]),
    ]
    for i, (title, icon, accent, items) in enumerate(columns):
        x = 0.55 + i * 2.55
        add_rect(slide, x, 2.0, 2.25, 4.45, "102D5E", "244879", rounded=True)
        add_icon_badge(slide, icon, x + 0.75, 2.28, 0.72, accent, WHITE, 13)
        add_text(slide, title, x + 0.25, 3.17, 1.75, 0.28, 12, accent, True, align=PP_ALIGN.CENTER)
        for j, item in enumerate(items):
            add_rect(slide, x + 0.2, 3.68 + j * 0.4, 1.85, 0.3, "173A70", "315990", rounded=True)
            add_text(slide, item, x + 0.28, 3.74 + j * 0.4, 1.68, 0.17, 8.8, WHITE, True, align=PP_ALIGN.CENTER)

    # 13. Data model and flow
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Data architecture", "One application record becomes the audit spine", "Structured entities and stored workflow state keep candidate evidence, decisions, scores, and media connected job-by-job.", number=13)
    entities = [
        ("USERS", "Identity, role, status", 0.55, 2.0, BLUE),
        ("PROFILES", "Candidate + recruiter details", 2.98, 2.0, CYAN),
        ("JOBS", "Details + assessment config", 5.41, 2.0, TEAL),
        ("APPLICATIONS", "Status + workflow + scores", 7.84, 2.0, VIOLET),
        ("INTERVIEWS", "Schedule + report + media", 10.27, 2.0, CORAL),
    ]
    for title, body, x, y, accent in entities:
        add_rect(slide, x, y, 2.08, 1.0, WHITE, LINE, rounded=True)
        add_text(slide, title, x + 0.15, y + 0.16, 1.78, 0.2, 10, accent, True, align=PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.15, y + 0.51, 1.78, 0.28, 9, SLATE, align=PP_ALIGN.CENTER)
    for x in [2.63, 5.06, 7.49, 9.92]:
        add_line(slide, x, 2.5, x + 0.35, 2.5, NAVY_2, 2)
    add_rect(slide, 1.05, 3.72, 11.2, 2.35, NAVY, NAVY, rounded=True)
    add_text(slide, "APPLICATION WORKFLOW STATE", 1.42, 4.02, 2.7, 0.25, 11, CYAN, True)
    flow_items = [
        ("Screening", "Approval / rejection", BLUE_2),
        ("Assessments", "Questions + answers", AMBER),
        ("Interview", "Turns + transcript", CORAL),
        ("Media", "Recording + snapshot", TEAL),
        ("Scoring", "ATS + final formula", VIOLET),
        ("Audit", "Events + timestamps", CYAN),
    ]
    for i, (label, body, accent) in enumerate(flow_items):
        x = 1.42 + i * 1.72
        add_circle(slide, x, 4.58, 0.32, accent)
        add_text(slide, label, x + 0.44, 4.52, 1.1, 0.22, 10, WHITE, True)
        add_text(slide, body, x + 0.44, 4.84, 1.05, 0.36, 8.3, "C4D8F7")
    add_text(slide, "Result: recruiters can open one candidate/job record and inspect the CV, workflow, scoring calculation, interview answers, media, and decision status.", 1.42, 5.5, 10.35, 0.32, 10.5, WHITE, True, align=PP_ALIGN.CENTER)

    # 14. Secure interview
    slide = prs.slides.add_slide(blank)
    add_background(slide, dark=True)
    add_title(slide, "Secure assessment and interview", "A controlled browser workflow with persistent evidence", "Security controls improve integrity and traceability while preserving a clear, guided candidate experience.", dark=True, number=14)
    stages = [
        ("01", "Device readiness", "Camera and microphone permission", BLUE_2),
        ("02", "Face presence", "One-face monitoring where supported", CYAN),
        ("03", "Verification", "Snapshot plus 10-second mic sample", TEAL),
        ("04", "Secure session", "Fullscreen and visibility warnings", AMBER),
        ("05", "Voice loop", "Question speech and live transcript", VIOLET),
        ("06", "Evidence save", "Video, audio, photo, transcript, report", CORAL),
    ]
    for i, (num, title, body, accent) in enumerate(stages):
        col = i % 3
        row = i // 3
        add_process_step(slide, num, title, body, 0.72 + col * 4.17, 2.05 + row * 1.72, 3.78, accent, dark=True)
    add_rect(slide, 0.72, 5.75, 11.85, 0.82, "173A70", "315990", rounded=True)
    add_text(slide, "Integrity policy", 0.98, 5.94, 1.35, 0.24, 11, CYAN, True)
    add_text(slide, "Leaving fullscreen or closing the secure terminal triggers warnings and can submit or lock the current attempt. Proctoring events are stored with the workflow.", 2.4, 5.92, 9.75, 0.34, 10.5, WHITE, True)

    # 15. Job ATS math
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Explainable scoring I", "Job ATS: evidence against one exact role", "The language model classifies evidence; the backend applies a deterministic, versioned formula. Grammar and visual formatting do not affect this score.", number=15)
    add_formula_box(slide, "Job ATS = 0.50R + 0.20S + 0.20E + 0.10D", 0.72, 1.95, 5.5, 0.78, NAVY, CYAN)
    add_text(slide, "R = requirement evidence   S = skill/tool coverage   E = relevant experience   D = responsibility readiness", 0.82, 2.86, 5.25, 0.45, 10, MUTED)
    components = [
        ("Requirement evidence", 50, BLUE, "Matched=1.00 | Partial=0.45 | Missing=0.00"),
        ("Skill and tool coverage", 20, CYAN, "Role-specific terms found in CV/profile evidence"),
        ("Experience alignment", 20, VIOLET, "40% tenure sufficiency + 60% domain/title fit"),
        ("Responsibility readiness", 10, TEAL, "Evidence for expected execution responsibilities"),
    ]
    for i, (label, weight, accent, note) in enumerate(components):
        y = 3.52 + i * 0.73
        add_text(slide, label, 0.78, y, 2.45, 0.22, 10.5, INK, True)
        add_rect(slide, 3.18, y + 0.03, 2.55, 0.16, "E5EDF8", "E5EDF8", rounded=True)
        add_rect(slide, 3.18, y + 0.03, 2.55 * weight / 50, 0.16, accent, accent, rounded=True)
        add_text(slide, f"{weight}%", 5.82, y - 0.02, 0.4, 0.22, 10, accent, True, align=PP_ALIGN.RIGHT)
        add_text(slide, note, 0.78, y + 0.29, 5.4, 0.22, 8.5, MUTED)
    add_rect(slide, 6.62, 1.95, 5.96, 4.76, WHITE, LINE, rounded=True)
    add_text(slide, "Requirement evidence weighting", 6.95, 2.23, 2.9, 0.3, 16, INK, True, TITLE_FONT)
    weight_items = [
        ("MANDATORY", "1.35x", CORAL, "Must / required / essential"),
        ("STANDARD", "1.00x", BLUE, "Normal role requirement"),
        ("PREFERRED", "0.70x", AMBER, "Nice-to-have / desirable"),
    ]
    for i, (label, value, accent, note) in enumerate(weight_items):
        y = 2.84 + i * 0.77
        add_rect(slide, 6.95, y, 5.28, 0.58, LIGHT, LINE, rounded=True)
        add_text(slide, label, 7.15, y + 0.13, 1.15, 0.2, 9, accent, True)
        add_text(slide, value, 8.38, y + 0.08, 0.75, 0.28, 15, accent, True, TITLE_FONT)
        add_text(slide, note, 9.28, y + 0.13, 2.62, 0.22, 9.5, SLATE)
    add_text(slide, "No artificial floor", 6.95, 5.42, 2.0, 0.28, 14, CORAL, True, TITLE_FONT)
    add_text(slide, "An unrelated CV can score near 0 instead of being forced to 35%. Each report also includes exact evidence, gaps, confidence, and formula components.", 6.95, 5.84, 5.15, 0.56, 10.5, SLATE)

    # 16. Final score math
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Explainable scoring II", "Final Score: job fit plus enabled assessment evidence", "Disabled or skipped rounds are excluded. Remaining weights are normalized, and incomplete workflows remain explicitly provisional.", number=16)
    add_formula_box(slide, "Final = weighted average(Job ATS, Aptitude, Technical, Coding, Interview)", 0.72, 1.92, 7.18, 0.7, NAVY, TEAL)
    score_example = [
        ("Job ATS", 80, 30, 24.0, BLUE),
        ("Aptitude", 100, 10, 10.0, CYAN),
        ("Technical MCQ", 80, 15, 12.0, VIOLET),
        ("Coding", 70, 20, 14.0, AMBER),
        ("Interview", 90, 25, 22.5, CORAL),
    ]
    for i, values in enumerate(score_example):
        label, score, weight, contribution, accent = values
        add_weight_bar(
            slide,
            label,
            score,
            weight,
            contribution,
            0.8,
            2.92 + i * 0.68,
            5.55,
            accent,
        )
    add_rect(slide, 7.52, 2.0, 5.0, 4.62, WHITE, LINE, rounded=True)
    add_text(slide, "Worked example", 7.88, 2.31, 2.1, 0.33, 18, INK, True, TITLE_FONT)
    add_text(slide, "(80x0.30) + (100x0.10) + (80x0.15) + (70x0.20) + (90x0.25)", 7.88, 2.93, 4.24, 0.58, 14, NAVY_2, True, font="Aptos Mono", align=PP_ALIGN.CENTER)
    add_text(slide, "82.5", 8.82, 3.82, 2.35, 0.8, 42, TEAL, True, TITLE_FONT, align=PP_ALIGN.CENTER)
    add_text(slide, "FINAL SCORE", 8.85, 4.55, 2.3, 0.25, 10, MUTED, True, align=PP_ALIGN.CENTER)
    add_rect(slide, 7.9, 5.16, 4.22, 0.95, PALE_BLUE, "C7DBFA", rounded=True)
    add_text(slide, "If Technical and Coding are disabled, their weights are removed and Job ATS + Aptitude + Interview weights are normalized to 100%.", 8.12, 5.36, 3.78, 0.48, 9.5, SLATE, True, align=PP_ALIGN.CENTER)

    # 17. Score usefulness
    slide = prs.slides.add_slide(blank)
    add_background(slide, dark=True)
    add_title(slide, "Decision usefulness", "Each score answers a different recruiter question", "Keeping profile quality, job fit, stage performance, and confidence separate prevents misleading comparisons.", dark=True, number=17)
    score_cards = [
        ("PROFILE ATS", "Is the CV/profile sufficiently complete and readable?", "General readiness", BLUE_2),
        ("MATCH SCORE", "Do declared skills overlap with posted requirements?", "Fast preview", CYAN),
        ("JOB ATS", "What evidence proves fit for this exact role?", "Screening priority", VIOLET),
        ("STAGE SCORES", "How did the candidate perform in each enabled round?", "Capability evidence", AMBER),
        ("FINAL SCORE", "What is the normalized outcome across available evidence?", "Overall comparison", TEAL),
        ("CONFIDENCE", "How strong and well-structured is the underlying evidence?", "Manual review signal", CORAL),
    ]
    for i, (label, question, use, accent) in enumerate(score_cards):
        col = i % 3
        row = i // 3
        x = 0.72 + col * 4.17
        y = 2.05 + row * 2.02
        add_rect(slide, x, y, 3.78, 1.65, "102D5E", "244879", rounded=True)
        add_text(slide, label, x + 0.22, y + 0.2, 1.45, 0.22, 9.5, accent, True)
        add_text(slide, question, x + 0.22, y + 0.58, 3.32, 0.45, 12, WHITE, True, TITLE_FONT)
        add_badge(slide, use, x + 0.22, y + 1.18, 1.65, "173A70", accent, "315990", 8)
    add_text(slide, "Scores support decisions; evidence and recruiter judgment remain the final authority.", 0.75, 6.36, 11.8, 0.28, 11, "DCEAFF", True, align=PP_ALIGN.CENTER)

    # 18. Analytics
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Recruiter intelligence", "Job-wise comparison instead of profile-only ranking", "Recruiters can move from portfolio KPIs to job-level demand, candidate evidence, performance, risk, recordings, and approval actions.", number=18)
    add_picture_crop(slide, SHOT_DIR / "recruiter-analytics.png", 0.68, 1.95, 8.15, 4.9)
    intelligence = [
        ("RANK", "Top 5 / 10 / 20 candidates by job", BLUE),
        ("COMPARE", "ATS, stages, final score, evidence", VIOLET),
        ("INSPECT", "CV, interview turns, transcript, media", CORAL),
        ("ACT", "Approve, reject, schedule, export", TEAL),
    ]
    for i, (label, body, accent) in enumerate(intelligence):
        y = 2.03 + i * 1.1
        add_rect(slide, 9.18, y, 3.35, 0.88, WHITE, LINE, rounded=True)
        add_badge(slide, label, 9.38, y + 0.16, 0.92, accent, WHITE, accent, 8)
        add_text(slide, body, 10.48, y + 0.19, 1.78, 0.42, 9.5, SLATE, True)
    add_formula_box(slide, "THE RECRUITER SEES THE WHY, NOT ONLY THE RANK", 9.18, 6.18, 3.35, 0.58, NAVY, CYAN)

    # 19. Smart automation
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Smart automation", "Convert unstructured hiring inputs into reviewable workflows", "Automation accelerates setup while preserving human confirmation before publishing or making decisions.", number=19)
    add_picture_crop(slide, SHOT_DIR / "recruiter-job-form.png", 0.68, 1.95, 6.72, 4.75)
    automation = [
        ("JOB AUTOFILL", "Paste text or upload PDF, Word, or image. Structured JSON fills title, type, experience, salary, requirements, responsibilities, and benefits.", BLUE),
        ("PROFILE AUTOFILL", "CV extraction fills candidate details, experience, skills, summary, links, and preferences with uncertainty warnings.", CYAN),
        ("QUESTION GENERATION", "Role-context aptitude, technical, coding, and interview questions vary between sessions.", VIOLET),
        ("HUMAN REVIEW", "Autofill never publishes automatically. Recruiters review values and warnings before submission.", TEAL),
    ]
    for i, (title, body, accent) in enumerate(automation):
        y = 1.98 + i * 1.18
        add_rect(slide, 7.76, y, 4.78, 0.97, WHITE, LINE, rounded=True)
        add_rect(slide, 7.76, y, 0.08, 0.97, accent, accent, rounded=True)
        add_text(slide, title, 8.05, y + 0.14, 1.65, 0.2, 9.5, accent, True)
        add_text(slide, body, 9.62, y + 0.13, 2.62, 0.62, 8.8, SLATE)
    add_badge(slide, "Low-temperature structured output + schema validation + evidence checks", 7.78, 6.2, 4.72, PALE_CYAN, NAVY_2, "BDEAF3", 8.5)

    # 20. Feasibility
    slide = prs.slides.add_slide(blank)
    add_background(slide, dark=True)
    add_title(slide, "Implementation feasibility", "Built with deployable technology and practical controls", "The current system proves technical and operational feasibility; enterprise hardening can be added incrementally without replacing the core architecture.", dark=True, number=20)
    feasibility = [
        ("TECHNICAL", "Mature web frameworks, REST APIs, browser media primitives, structured model outputs, and managed PostgreSQL.", BLUE_2, "PROVEN"),
        ("OPERATIONAL", "Role-specific interfaces, guided stage gates, review-before-publish, immutable decisions, and exports.", TEAL, "USABLE"),
        ("ECONOMIC", "Open-source application stack with usage-based external intelligence and container hosting.", AMBER, "LEAN"),
        ("SCALABILITY", "Separated services, stateless API patterns, pooled DB connectivity, and persistent storage volumes.", VIOLET, "READY"),
    ]
    for i, (title, body, accent, state) in enumerate(feasibility):
        x = 0.72 + (i % 2) * 6.05
        y = 2.05 + (i // 2) * 1.72
        add_rect(slide, x, y, 5.65, 1.38, "102D5E", "244879", rounded=True)
        add_text(slide, title, x + 0.23, y + 0.2, 1.52, 0.22, 10, accent, True)
        add_badge(slide, state, x + 4.25, y + 0.16, 1.05, "173A70", accent, "315990", 8)
        add_text(slide, body, x + 0.23, y + 0.62, 5.05, 0.48, 10.2, WHITE)
    add_rect(slide, 0.72, 5.57, 11.7, 0.83, "173A70", "315990", rounded=True)
    add_text(slide, "CURRENT PRODUCTION", 0.98, 5.8, 1.55, 0.2, 9, CYAN, True)
    add_text(slide, "Frontend + backend on Railway | PostgreSQL service | persistent CV/interview media volume | health checks | environment-managed secrets", 2.6, 5.74, 9.35, 0.34, 10.5, WHITE, True)

    # 21. Security and governance
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Security and governance", "Defense-in-depth across identity, workflow, and media", "The implementation combines preventive controls, evidence capture, and role-separated operations.", number=21)
    layers = [
        ("IDENTITY", ["Password hashing", "JWT bearer tokens", "Active-account checks", "Role-specific route guards"], BLUE),
        ("WORKFLOW", ["Recruiter approval gate", "No duplicate application", "Irreversible approve/reject", "Disabled stages omitted"], VIOLET),
        ("ASSESSMENT", ["Fullscreen events", "Visibility warnings", "Camera and mic readiness", "Face-presence signal"], AMBER),
        ("EVIDENCE", ["Snapshot", "Video + audio recording", "Transcript + question turns", "Timestamped report"], CORAL),
        ("PLATFORM", ["CORS allowlist/regex", "Environment secrets", "Upload size controls", "Persistent storage"], TEAL),
    ]
    for i, (title, items, accent) in enumerate(layers):
        x = 0.55 + i * 2.55
        add_rect(slide, x, 2.05, 2.25, 3.8, WHITE, LINE, rounded=True)
        add_icon_badge(slide, str(i + 1), x + 0.77, 2.35, 0.7, accent, WHITE, 15)
        add_text(slide, title, x + 0.25, 3.27, 1.75, 0.24, 10, accent, True, align=PP_ALIGN.CENTER)
        add_bullets(slide, items, x + 0.23, 3.77, 1.8, 1.7, 9.2, SLATE, accent, 4)
    add_formula_box(slide, "Enterprise next step: SSO/SAML + RBAC policies + encrypted object storage + retention controls + audit export", 0.72, 6.18, 11.85, 0.58, NAVY, CORAL)

    # 22. Competitive positioning
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Competitive positioning", "Why Invenia AI is differentiated", "The USP is not one AI feature; it is the connected, explainable decision workflow across screening, assessment, interview, and analytics.", number=22)
    headers = ["CAPABILITY", "TRADITIONAL ATS", "TEST PLATFORM", "VIDEO TOOL", "INVENIA AI"]
    widths = [3.05, 2.15, 2.05, 2.0, 2.4]
    xs = [0.45]
    for width in widths[:-1]:
        xs.append(xs[-1] + width)
    y0 = 1.95
    for i, header in enumerate(headers):
        fill = NAVY if i == 4 else NAVY_2
        add_rect(slide, xs[i], y0, widths[i], 0.55, fill, fill, rounded=(i in {0, 4}))
        add_text(slide, header, xs[i] + 0.08, y0 + 0.15, widths[i] - 0.16, 0.2, 8.5, WHITE, True, align=PP_ALIGN.CENTER)
    rows = [
        ("Job-context CV evidence", "Limited", "No", "No", "Strong"),
        ("Configurable stage workflow", "Partial", "Partial", "No", "End-to-end"),
        ("Secure voice interview", "No", "No", "Video only", "Integrated"),
        ("Deterministic score formula", "Opaque", "Test only", "No", "Explainable"),
        ("Interview media + transcript", "External", "External", "Yes", "Connected"),
        ("Job-wise comparison + export", "Basic", "Partial", "No", "Unified"),
    ]
    for r, row in enumerate(rows):
        y = y0 + 0.58 + r * 0.62
        for c, value in enumerate(row):
            fill = PALE_BLUE if c == 4 else (WHITE if r % 2 == 0 else "F7FAFF")
            add_rect(slide, xs[c], y, widths[c], 0.58, fill, LINE)
            color = TEAL if c == 4 else (INK if c == 0 else SLATE)
            add_text(slide, value, xs[c] + 0.1, y + 0.16, widths[c] - 0.2, 0.2, 9.2, color, c in {0, 4}, align=PP_ALIGN.CENTER if c else PP_ALIGN.LEFT)
    usps = [
        ("01", "Evidence-grounded", "Scores show matched, partial, missing, and quoted evidence", BLUE),
        ("02", "Workflow-aware", "Only recruiter-enabled stages influence candidate flow and final score", VIOLET),
        ("03", "Media-complete", "CV, snapshot, recording, transcript, report, and exports stay connected", CORAL),
    ]
    for i, (num, title, body, accent) in enumerate(usps):
        add_card(slide, title, body, 0.7 + i * 4.18, 6.05, 3.8, 0.77, num, accent, WHITE, LINE, 12, 8.5)

    # 23. Future scope
    slide = prs.slides.add_slide(blank)
    add_background(slide, dark=True)
    add_title(slide, "Future scope", "Roadmap from commercial MVP to enterprise talent platform", "The architecture supports incremental hardening, deeper automation, and enterprise integrations.", dark=True, number=23)
    roadmap = [
        ("NOW", "Commercial MVP", ["Dual portals", "Grounded ATS", "Assessments", "Recorded interview", "Analytics + export"], BLUE_2),
        ("NEXT", "Enterprise readiness", ["SSO/SAML", "Fine-grained RBAC", "Object storage/CDN", "Audit retention", "Automated backups"], CYAN),
        ("SCALE", "Hiring ecosystem", ["HRIS/ATS connectors", "Calendar + email", "Webhook/API marketplace", "Multilingual interviews", "Mobile candidate app"], TEAL),
        ("INTELLIGENCE", "Advanced decision science", ["Validated benchmark sets", "Bias/fairness audits", "Skills ontology", "Talent rediscovery", "Predictive workforce insights"], VIOLET),
    ]
    for i, (phase, title, items, accent) in enumerate(roadmap):
        x = 0.62 + i * 3.13
        add_rect(slide, x, 2.0, 2.85, 4.35, "102D5E", "244879", rounded=True)
        add_badge(slide, phase, x + 0.22, 2.28, 0.86, "173A70", accent, "315990", 8)
        add_text(slide, title, x + 0.22, 2.85, 2.4, 0.55, 18, WHITE, True, TITLE_FONT)
        add_bullets(slide, items, x + 0.22, 3.62, 2.4, 2.15, 10, "DCEAFF", accent, 5)
        add_text(slide, f"0{i + 1}", x + 2.18, 5.78, 0.42, 0.3, 16, accent, True, TITLE_FONT, align=PP_ALIGN.RIGHT)

    # 24. Close
    slide = prs.slides.add_slide(blank)
    add_background(slide, dark=True)
    add_rect(slide, 0.68, 0.58, 1.38, 0.94, WHITE, WHITE, rounded=True)
    slide.shapes.add_picture(str(LOGO), Inches(0.79), Inches(0.68), width=Inches(1.16), height=Inches(0.84))
    add_text(slide, "INVENIA AI", 0.72, 2.04, 7.4, 0.72, 38, WHITE, True, TITLE_FONT)
    add_text(slide, "Hire with evidence. Assess with integrity. Decide with confidence.", 0.74, 2.91, 9.8, 0.48, 21, CYAN, True, TITLE_FONT)
    add_text(slide, "A live, full-stack recruitment platform for the next generation of enterprise hiring.", 0.74, 3.65, 8.2, 0.45, 14, "DCEAFF")
    add_rect(slide, 8.83, 1.63, 3.65, 3.85, "102D5E", "29548C", rounded=True)
    add_text(slide, "EXHIBITION DEMO", 9.18, 2.02, 2.94, 0.24, 10, CYAN, True, align=PP_ALIGN.CENTER)
    demo_items = ["Role-based login", "Smart job creation", "Candidate job ATS", "Approval-gated rounds", "Recruiter analytics"]
    for i, item in enumerate(demo_items):
        add_icon_badge(slide, str(i + 1), 9.22, 2.52 + i * 0.49, 0.34, [BLUE_2, CYAN, VIOLET, AMBER, TEAL][i], WHITE, 7.5)
        add_text(slide, item, 9.72, 2.55 + i * 0.49, 2.13, 0.2, 10, WHITE, True)
    link = add_text(slide, "OPEN LIVE PRODUCT", 9.35, 5.0, 2.55, 0.27, 10, CYAN, True, align=PP_ALIGN.CENTER)
    link.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://frontend-production-fad4.up.railway.app/"
    add_text(slide, "Thank you", 0.74, 5.25, 3.5, 0.5, 24, WHITE, True, TITLE_FONT)
    add_text(slide, "Questions and product walkthrough", 0.74, 5.83, 4.7, 0.32, 12, "AFC8EF")
    add_footer(slide, 24, dark=True)

    return prs


def validate_deck(prs: Presentation) -> list[str]:
    issues: list[str] = []
    if len(prs.slides) != 24:
        issues.append(f"Expected 24 slides, found {len(prs.slides)}")
    for slide_index, slide in enumerate(prs.slides, start=1):
        if len(slide.shapes) < 3:
            issues.append(f"Slide {slide_index} has too few shapes")
        for shape in slide.shapes:
            left = shape.left / Inches(1)
            top = shape.top / Inches(1)
            right = left + shape.width / Inches(1)
            bottom = top + shape.height / Inches(1)
            # Large translucent circles intentionally bleed beyond the slide edge.
            # Keep a generous artboard guard while still catching accidental drift.
            if left < -2.05 or top < -1.9 or right > SLIDE_W + 1.75 or bottom > SLIDE_H + 1.4:
                issues.append(
                    f"Slide {slide_index} shape out of bounds: ({left:.2f}, {top:.2f}, {right:.2f}, {bottom:.2f})"
                )
    return issues


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    required = [
        LOGO,
        SHOT_DIR / "candidate-jobs.png",
        SHOT_DIR / "candidate-profile.png",
        SHOT_DIR / "recruiter-overview.png",
        SHOT_DIR / "recruiter-analytics.png",
        SHOT_DIR / "recruiter-job-form.png",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing presentation assets:\n" + "\n".join(missing))
    prs = create_deck()
    issues = validate_deck(prs)
    if issues:
        raise RuntimeError("Presentation validation failed:\n" + "\n".join(issues))
    prs.save(OUT_PATH)
    print(f"Created {OUT_PATH}")
    print(f"Slides: {len(prs.slides)}")
    print(f"Size: {OUT_PATH.stat().st_size} bytes")


if __name__ == "__main__":
    main()
